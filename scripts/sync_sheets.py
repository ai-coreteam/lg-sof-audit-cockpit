#!/usr/bin/env python3
"""
LG Electronics - SOF Monitoring & Governance Cockpit
Automated Ingestion & Real-Time Sync Pipeline for Google Sheet (W38-W42 Middle Live)
Target URL: https://docs.google.com/spreadsheets/d/1lA6mxuh3yn6TMGWH8TIWOl6PWyKP5PeZ586sOnOGG40/edit?gid=1357809752#gid=1357809752
"""

import os
import sys
import csv
import io
import json
import zipfile
import urllib.request
import xml.etree.ElementTree as ET

GOOGLE_SHEET_ID = "1lA6mxuh3yn6TMGWH8TIWOl6PWyKP5PeZ586sOnOGG40"
SHEET_GID = "1357809752"
CSV_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/export?format=csv&gid={SHEET_GID}"
GVIZ_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/gviz/tq?tqx=out:json&sheet=Check%20theo%20CAT,%20theo%20Brand"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/export?format=xlsx"

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(PROJECT_ROOT, "data", "google_sheet_latest.xlsx")
W41_JSON_PATH = os.path.join(PROJECT_ROOT, "data", "w41_audit_cockpit.json")
OUTPUT_JSON_W42 = os.path.join(PROJECT_ROOT, "data", "w42_audit_cockpit.json")
OUTPUT_JS_W42 = os.path.join(PROJECT_ROOT, "app", "assets", "w42_audit_cockpit.js")
ROOT_ASSETS_JS_W42 = os.path.join(PROJECT_ROOT, "assets", "w42_audit_cockpit.js")


def classify_dealer(store_name):
    c_upper = str(store_name or "").upper()
    if "DML" in c_upper or "DMM" in c_upper or "DIEN MAY XANH" in c_upper:
        return "DMX"
    if "CP" in c_upper or "CAO PHONG" in c_upper or "CAOPHONG" in c_upper or "CHO LON" in c_upper:
        return "CP"
    if "MM" in c_upper or "MEGA" in c_upper:
        return "MM"
    if "VHC" in c_upper or "HC" in c_upper:
        return "VHC"
    return "Tradi"


def load_store_master():
    store_map = {}
    if not os.path.exists(XLSX_PATH):
        return store_map

    try:
        import openpyxl
        wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)
        sheet_name = "Mapping PC SUP" if "Mapping PC SUP" in wb.sheetnames else "Sheet6"
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for idx, row in enumerate(ws.iter_rows(values_only=True)):
                if idx > 1 and row and row[0]:
                    code = str(row[0]).strip()
                    dealer = classify_dealer(row[1])
                    store_map[code] = {
                        "store_name": row[1] or "",
                        "region": row[2] or "",
                        "province": row[3] or "",
                        "district": row[4] or "",
                        "sup": row[5] or "",
                        "dealer": dealer,
                        "channel": "MC" if dealer in ["DMX", "CP", "MM", "VHC"] else "Tradi",
                        "active_pc": int(float(row[11])) if len(row) > 11 and row[11] is not None else 0
                    }
        return store_map
    except Exception:
        pass

    # Pure Python stdlib fallback (zipfile + xml.etree.ElementTree)
    try:
        with zipfile.ZipFile(XLSX_PATH) as z:
            ss_xml = ET.fromstring(z.read("xl/sharedStrings.xml"))
            ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            shared_strings = ["".join(t.text or "" for t in si.findall(".//a:t", ns)) for si in ss_xml.findall("a:si", ns)]
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            rels_xml = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            rels = {r.attrib["Id"]: r.attrib["Target"] for r in rels_xml}
            sheet_target = None
            for s in wb_xml.findall(".//a:sheet", ns):
                if s.attrib["name"] in ("Mapping PC SUP", "Sheet6"):
                    rid = s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                    sheet_target = "xl/" + rels[rid]
            if sheet_target:
                ws_xml = ET.fromstring(z.read(sheet_target))
                for row in ws_xml.findall(".//a:row", ns)[2:]:
                    vals = {}
                    for c in row.findall("a:c", ns):
                        ref = "".join(ch for ch in c.attrib["r"] if ch.isalpha())
                        t = c.attrib.get("t")
                        v_el = c.find("a:v", ns)
                        if v_el is not None and v_el.text is not None:
                            vals[ref] = shared_strings[int(v_el.text)] if t == "s" else v_el.text
                    code = vals.get("A", "").strip()
                    if code:
                        sname = vals.get("B", "")
                        dealer = classify_dealer(sname)
                        store_map[code] = {
                            "store_name": sname,
                            "region": vals.get("C", ""),
                            "province": vals.get("D", ""),
                            "district": vals.get("E", ""),
                            "sup": vals.get("F", ""),
                            "dealer": dealer,
                            "channel": "MC" if dealer in ["DMX", "CP", "MM", "VHC"] else "Tradi",
                            "active_pc": int(float(vals.get("L", 0))) if vals.get("L") else 0
                        }
    except Exception as e:
        print("Notice: cannot read XLSX store master:", e)
    return store_map


def load_w41_hints():
    exact_hints = {}
    shop_hints = {}
    if os.path.exists(W41_JSON_PATH):
        try:
            with open(W41_JSON_PATH, "r", encoding="utf-8") as f:
                w41_data = json.load(f)
            for r in w41_data.get("records", []):
                rec_id = r.get("id", "")
                scode = r.get("store_code", "")
                cat = r.get("category", "")
                brand = r.get("brand", "")
                extra = (r.get("extra_note") or "").strip()
                fb = (r.get("feedback_raw") or "").strip()
                note_summary = f"Extra Note W41: \"{extra}\"" if extra else (f"Feedback W41: \"{fb}\"" if fb else "")
                if note_summary:
                    exact_hints[rec_id] = f"Ca lặp lại từ W41 ({cat} {brand}) — {note_summary}"
                    shop_hints.setdefault(scode, []).append(f"{cat} {brand} ({r.get('w40', 0)}→{r.get('w41', 0)}): {note_summary}")
        except Exception as e:
            print("Notice: cannot read W41 hints:", e)
    return exact_hints, shop_hints


def evaluate_feedback(record):
    fb_raw = (record.get("feedback_raw") or "").strip()
    brand = record.get("brand", "")
    w_cur = record.get("w42", record.get("w41", 0))
    w_prev = record.get("w41", record.get("w40", 0)) if "w42" in record else record.get("w40", 0)
    delta = w_cur - w_prev
    flags = record.get("flags", {})

    if not fb_raw and not (record.get("extra_note") or "").strip():
        return {
            "status": "CHƯA_FEEDBACK",
            "validity": "PENDING",
            "validity_label": "Chưa có phản hồi",
            "validity_desc": "PC SUP chưa điền giải trình trên Google Sheet",
            "risk_level": "HIGH",
            "risk_label": "Rủi ro tồn đọng",
            "risk_desc": "Cần đôn đốc SUP phản hồi gấp trước hạn chốt tuần",
            "action_code": "PUSH_SUP",
            "action_title": "Đôn đốc phản hồi gấp",
            "action_detail": "Gửi tin nhắn cảnh báo yêu cầu SUP xác minh trong ngày."
        }

    fb_lower = fb_raw.lower()
    extra_note = (record.get("extra_note") or "").strip().replace("\r\n", " · ").replace("\n", " · ")
    extra_lower = extra_note.lower()

    # Rule SAN-03: Explicit PC error admission in Feedback dropdown (prioritized even if Extra Note is also attached)
    if "báo sai" in fb_lower or "nhập sai" in fb_lower or "điều chỉnh lại" in fb_lower:
        desc = (
            f"SUP thừa nhận PC báo sai tuần cũ; bổ sung thực tế: \"{extra_note}\"."
            if extra_note
            else "SUP thừa nhận nhân viên PC báo sai số tuần trước, tuần này sửa lại"
        )
        detail = (
            f"Kiểm điểm PC báo sai tuần cũ & cập nhật theo thực địa: '{extra_note}'."
            if extra_note
            else "Yêu cầu SUP đào tạo lại quy trình kiểm đếm cho PC tại điểm bán."
        )
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "PC_ERROR",
            "validity_label": "Thừa nhận PC báo sai",
            "validity_desc": desc,
            "risk_level": "P2_OPERATIONAL",
            "risk_label": "Lỗi vận hành PC",
            "risk_desc": "Số liệu quá khứ không đáng tin cậy. Cần kiểm điểm kỷ luật PC.",
            "action_code": "COACH_PC",
            "action_title": "Kiểm điểm & Đào tạo PC",
            "action_detail": detail
        }

    # Rule SAN-EXTRA: If SUP provided an Extra note, this case is EXCLUSIVELY clarified via Extra Note
    if extra_note:
        # Check if it's count correction / data mismatch (including "week trước chỉ 44set")
        if any(k in extra_lower for k in ["số đúng là", "không phải", "khổng phải", "sai số", "lệch số", "nhập nhầm", "week trước chỉ", "tuần trước chỉ", "check data"]):
            return {
                "status": "ĐÃ_FEEDBACK",
                "validity": "DATA_MISMATCH",
                "validity_label": "Sai lệch số thực tế",
                "validity_desc": f"SUP bổ sung: \"{extra_note}\". Số trên hệ thống không khớp với thực tế kiểm đếm.",
                "risk_level": "P2_OPERATIONAL",
                "risk_label": "Dữ liệu Sheet bị lệch số",
                "risk_desc": "Số liệu trên Sheet khai khống so với thực tế. Cần hiệu chỉnh lại số đúng.",
                "action_code": "CORRECT_DATA",
                "action_title": "Bác bỏ & Hiệu chỉnh số liệu",
                "action_detail": f"Yêu cầu sửa số liệu chính thức theo ghi chú thực tế: '{extra_note}'; kiểm điểm PC."
            }

        # Check if competitor shelf cut / delisting (excluding balanced model rotation "lên ... cắt ...")
        if any(k in extra_lower for k in ["cắt", "bỏ mẫu", "đảo"]) and "lên " not in extra_lower and brand != "LG":
            return {
                "status": "ĐÃ_FEEDBACK",
                "validity": "PLAUSIBLE",
                "validity_label": "Đã làm rõ (Extra Note)",
                "validity_desc": f"SUP bổ sung thông tin thực địa: \"{extra_note}\".",
                "risk_level": "P3_NORMAL",
                "risk_label": f"Đối thủ {brand} bị cắt giảm",
                "risk_desc": f"Thị trường thu hẹp trưng bày model đối thủ {brand} tại cửa hàng.",
                "action_code": "EXPAND_LG",
                "action_title": "Tận dụng chiếm đảo kệ",
                "action_detail": f"{brand} bị cắt giảm trưng bày. Đề xuất làm việc với shop để đưa thêm sản phẩm LG vào đảo trống."
            }

        # All other Extra Notes: Fully clarify the case and remove from SUSPICIOUS
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "PLAUSIBLE",
            "validity_label": "Đã làm rõ (Extra Note)",
            "validity_desc": f"SUP bổ sung: \"{extra_note}\". Đã làm rõ nguyên nhân biến động thực tế.",
            "risk_level": "P3_NORMAL",
            "risk_label": "Biến động đã làm rõ",
            "risk_desc": "Cửa hàng đã giải trình cụ thể biến động thực tế.",
            "action_code": "ACCEPT",
            "action_title": "Phê duyệt giải trình (Có Extra Note)",
            "action_detail": f"Chấp thuận do SUP đã bổ sung nguyên nhân thực địa: '{extra_note}'."
        }

    # Rule SAN-01: Data drops to 0 but supervisor claims "normal fluctuations"
    if flags.get("zero_drop") and "tăng giảm đúng thực tế" in fb_lower:
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "SUSPICIOUS",
            "validity_label": "Nghi vấn đối phó",
            "validity_desc": "Dữ liệu về 0 nhưng ghi tăng giảm đúng thực tế. Cần ảnh chụp quầy.",
            "risk_level": "P0_CRITICAL" if brand == "LG" else "P1_SUSPICIOUS",
            "risk_label": "Mất SOF LG" if brand == "LG" else "Nghi vấn đối phó",
            "risk_desc": f"Nguy cơ mất trắng trưng bày {brand} tại shop nhưng phản hồi qua loa",
            "action_code": "REJECT_AND_AUDIT",
            "action_title": "Bác bỏ & Đòi ảnh thực tế",
            "action_detail": "Từ chối duyệt. Yêu cầu SUP gửi ảnh chụp kệ thực tế trong 24h."
        }

    # Rule SAN-02: Radical swing >= 8 sets with superficial explanation
    if flags.get("diff_8_4w") and "tăng giảm đúng thực tế" in fb_lower:
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "SUSPICIOUS",
            "validity_label": "Nghi vấn đối phó",
            "validity_desc": "Biến động lệch >8 sets nhưng giải trình chung chung, thiếu căn cứ.",
            "risk_level": "P1_SUSPICIOUS",
            "risk_label": "Nghi vấn đối phó",
            "risk_desc": "Khả năng copy-paste giải trình hàng loạt không qua kiểm tra thực tế",
            "action_code": "REJECT_AND_AUDIT",
            "action_title": "Yêu cầu ảnh & Biên bản",
            "action_detail": "Yêu cầu SUP giải trình chi tiết lý do biến động đột biến."
        }

    # Commercial Risk: Delisting or shelf cut
    if any(k in fb_lower for k in ["cắt shop", "bỏ quầy", "dồn quầy", "hạ bục", "thu hồi"]):
        risk = "P0_CRITICAL" if brand == "LG" else "P1_COMMERCIAL"
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "PLAUSIBLE",
            "validity_label": "Hợp lý (Có căn cứ)",
            "validity_desc": "Giải trình phản ánh việc đại lý cắt giảm hoặc thu hồi quầy kệ",
            "risk_level": risk,
            "risk_label": "P0: Mất quầy LG" if brand == "LG" else "P1: Cắt quầy đại lý",
            "risk_desc": f"Đại lý đang thu hẹp mặt bằng trưng bày {brand}",
            "action_code": "ESCALATE_KAM",
            "action_title": "Can thiệp KAM khẩn cấp",
            "action_detail": "Báo cáo Head of GTM & làm việc với KAM để bảo vệ hợp đồng quầy kệ."
        }

    # LG Display drop risk
    if brand == "LG" and delta < 0 and abs(delta) >= 2:
        return {
            "status": "ĐÃ_FEEDBACK",
            "validity": "PLAUSIBLE",
            "validity_label": "Hợp lý đạt chuẩn",
            "validity_desc": f"LG giảm {abs(delta)} sets. Cần theo dõi để tránh mất thị phần SOF.",
            "risk_level": "P0_CRITICAL",
            "risk_label": "Rủi ro giảm trưng bày LG",
            "risk_desc": f"Số lượng trưng bày của LG bị giảm {abs(delta)} máy",
            "action_code": "ESCALATE_KAM",
            "action_title": "Theo dõi SOF & Kiểm tra shop",
            "action_detail": "Kiểm tra lý do giảm model LG và đôn đốc đưa thêm hàng trưng bày."
        }

    # Normal Plausible explanation
    return {
        "status": "ĐÃ_FEEDBACK",
        "validity": "PLAUSIBLE",
        "validity_label": "Hợp lý đạt chuẩn",
        "validity_desc": "Lý do phù hợp với biến động luân chuyển hàng hóa thực tế",
        "risk_level": "P3_NORMAL",
        "risk_label": "Biến động bình thường",
        "risk_desc": "Thay đổi số lượng do relayout hoặc luân chuyển model bình thường",
        "action_code": "ACCEPT",
        "action_title": "Phê duyệt giải trình",
        "action_detail": "Chấp thuận dữ liệu đưa vào báo cáo SOF chính thức."
    }


def fetch_live_csv(csv_file_path=None):
    if csv_file_path and os.path.exists(csv_file_path):
        print(f"[1/3] Reading CSV data from local snapshot {csv_file_path}...")
        with open(csv_file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        # Strip markdown frontmatter if present from tool cache
        if content.startswith("Title:") and "\n---\n" in content:
            content = content.split("\n---\n", 1)[1].lstrip("\r\n")
    else:
        print(f"[1/3] Fetching live data via Google Sheets CSV Export from {CSV_URL}...")
        req = urllib.request.Request(CSV_URL, headers={"User-Agent": "Mozilla/5.0"})
        content = urllib.request.urlopen(req).read().decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(content))
    return list(reader)


def detect_header_indices(header_row):
    """Self-healing header parser to handle swapped Region / PC Active columns across cycles."""
    idx_region = 3
    idx_pc = 4
    idx_extra = 17
    for idx, col_name in enumerate(header_row):
        c_low = str(col_name or "").strip().lower()
        if c_low == "region" or "khu vực" in c_low:
            idx_region = idx
        elif "pc" in c_low and "sup" not in c_low:
            idx_pc = idx
        elif "extra" in c_low or "ghi chú" in c_low:
            idx_extra = idx
    return idx_region, idx_pc, idx_extra


def sync_live(csv_file_path=None):
    store_master = load_store_master()
    exact_hints, shop_hints = load_w41_hints()
    rows = fetch_live_csv(csv_file_path)
    print(f"[2/3] Processing {len(rows)} rows from Google Sheets calculation engine...")

    header_row = rows[1] if len(rows) > 1 else []
    idx_region, idx_pc, idx_extra = detect_header_indices(header_row)

    records = []
    for i, r in enumerate(rows):
        if i <= 1 or len(r) < 16:
            continue

        d2_str = str(r[12] or "").strip().lower()
        d8_str = str(r[13] or "").strip().lower()
        z0_str = str(r[14] or "").strip().lower()

        is_d2 = "x" in d2_str or d2_str == "1"
        is_d8 = "x" in d8_str or d8_str == "1"
        is_z0 = "x" in z0_str or z0_str == "1"

        if is_d2 or is_d8 or is_z0:
            shop_code = str(r[0] or "").strip()
            master_info = store_master.get(shop_code, {})
            shop_name = r[1] or master_info.get("store_name", "")
            sup = r[2] or master_info.get("sup", "")

            reg_raw = str(r[idx_region] or "").strip() if len(r) > idx_region else ""
            reg = reg_raw if reg_raw in ("Middle", "North", "South") else (master_info.get("region") or "Middle")

            pc_act_raw = str(r[idx_pc] or "").strip() if len(r) > idx_pc else ""
            try:
                pc_act = int(float(pc_act_raw)) if pc_act_raw else master_info.get("active_pc", 0)
            except (ValueError, TypeError):
                pc_act = master_info.get("active_pc", 0)

            cat = str(r[5] or "").strip()
            brand = str(r[6] or "").strip()
            try:
                w38 = int(float(r[7])) if r[7] and str(r[7]).strip() else 0
                w39 = int(float(r[8])) if r[8] and str(r[8]).strip() else 0
                w40 = int(float(r[9])) if r[9] and str(r[9]).strip() else 0
                w41 = int(float(r[10])) if r[10] and str(r[10]).strip() else 0
                w42 = int(float(r[11])) if r[11] and str(r[11]).strip() else 0
            except (ValueError, TypeError):
                w38 = w39 = w40 = w41 = w42 = 0

            fb = str(r[15] or "").strip()
            extra_note = str(r[idx_extra] or "").strip() if len(r) > idx_extra else ""
            dealer = master_info.get("dealer") or classify_dealer(shop_name)
            channel = "MC" if dealer in ["DMX", "CP", "MM", "VHC"] else "Tradi"

            rec_id = f"{shop_code}_{cat}_{brand}"
            w41_hint = exact_hints.get(rec_id, "")
            if not w41_hint and shop_code in shop_hints:
                w41_hint = "Lịch sử shop W41 (ngành khác): " + " | ".join(shop_hints[shop_code])

            rec = {
                "id": rec_id,
                "store_code": shop_code,
                "store_name": shop_name,
                "dealer": dealer,
                "channel": channel,
                "region": reg,
                "province": master_info.get("province", ""),
                "district": master_info.get("district", ""),
                "sup": sup,
                "active_pc": pc_act,
                "category": cat,
                "brand": brand,
                "w38": w38, "w39": w39, "w40": w40, "w41": w41, "w42": w42,
                "delta_wow": w42 - w41,
                "flags": {
                    "diff_8_4w": is_d8,
                    "diff_2_prev": is_d2,
                    "zero_drop": is_z0,
                    "lg_drop_50": bool(brand == "LG" and w41 > 0 and w42 <= w41 * 0.5)
                },
                "feedback_raw": fb,
                "extra_note": extra_note,
                "w41_hint": w41_hint
            }
            rec["eval"] = evaluate_feedback(rec)
            records.append(rec)

    middle_store_master = {
        code: info for code, info in store_master.items()
        if info.get("region") == "Middle" or any(r["store_code"] == code for r in records)
    }

    data_payload = {
        "audit_cycle": "W38-W42 (Google Sheet Live)",
        "cycle_short": "W42",
        "google_sheet_url": f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/edit?gid={SHEET_GID}#gid={SHEET_GID}",
        "weeks": ["W38", "W39", "W40", "W41", "W42"],
        "total_records": len(records),
        "records": records,
        "w41_exact_hints": exact_hints,
        "w41_shop_hints": shop_hints,
        "store_master": middle_store_master
    }

    print(f"[3/3] Saving {len(records)} live audited W42 records...")
    with open(OUTPUT_JSON_W42, "w", encoding="utf-8") as f:
        json.dump(data_payload, f, ensure_ascii=False, indent=2)

    with open(OUTPUT_JS_W42, "w", encoding="utf-8") as f:
        f.write("// LG Electronics SOF Governance Cockpit - Live Google Sheet W42 Dataset\n")
        f.write("window.SOF_AUDIT_DATA_W42 = ")
        json.dump(data_payload, f, ensure_ascii=False, indent=2)
        f.write(";\n")

    print(f"      Saved JSON: {OUTPUT_JSON_W42}")
    print(f"      Saved JS: {OUTPUT_JS_W42}")

    if os.path.exists(os.path.dirname(ROOT_ASSETS_JS_W42)):
        with open(ROOT_ASSETS_JS_W42, "w", encoding="utf-8") as f:
            f.write("// LG Electronics SOF Governance Cockpit - Live Google Sheet W42 Dataset\n")
            f.write("window.SOF_AUDIT_DATA_W42 = ")
            json.dump(data_payload, f, ensure_ascii=False, indent=2)
            f.write(";\n")
        print(f"      Saved Root JS: {ROOT_ASSETS_JS_W42}")

    print("SUCCESS: Real-time W42 sync completed successfully!")


if __name__ == "__main__":
    csv_arg = sys.argv[1] if len(sys.argv) > 1 else None
    sync_live(csv_arg)

